from __future__ import annotations

import hashlib
import re
from datetime import date, datetime, time
from decimal import Decimal
from collections import defaultdict
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from offline_masker.domain.enums import DEFAULT_METHODS, SensitiveType
from offline_masker.domain.models import Candidate, ColumnRule, CustomTerm, ScanResult
from offline_masker.security.safe_preview import safe_preview
from offline_masker.security.hashing import value_kind

from .custom_term_detector import find_custom_term
from .header_detector import infer_header_type
from .regex_detector import detect_regex


WHOLE_CELL_COLUMN_TYPES = {
    SensitiveType.BANK_ACCOUNT,
    SensitiveType.PERSON_NAME,
    SensitiveType.CUSTOMER,
    SensitiveType.SUPPLIER,
    SensitiveType.COMPANY,
    SensitiveType.ADDRESS,
    SensitiveType.CONTRACT,
}


def infer_column_rules(
    data: bytes,
    sheet_names: list[str],
    enabled_types: set[SensitiveType],
    header_row: int = 1,
) -> list[ColumnRule]:
    workbook = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=True, rich_text=True)
    rules: list[ColumnRule] = []
    try:
        for sheet_name in sheet_names:
            worksheet = workbook[sheet_name]
            for column_index in range(1, worksheet.max_column + 1):
                inferred = infer_header_type(worksheet.cell(header_row, column_index).value)
                if inferred and inferred[0] in enabled_types:
                    rules.append(
                        ColumnRule(
                            sheet=sheet_name,
                            column=get_column_letter(column_index),
                            sensitive_type=inferred[0],
                            header_row=header_row,
                            source=inferred[1],
                        )
                    )
    finally:
        workbook.close()
    return rules


class DetectionCoordinator:
    def scan(
        self,
        data: bytes,
        sheet_names: list[str],
        enabled_types: set[SensitiveType],
        column_rules: list[ColumnRule] | None = None,
        custom_terms: list[CustomTerm] | None = None,
        *,
        restrict_to_column_rules: bool = False,
    ) -> ScanResult:
        column_rules = column_rules or []
        custom_terms = custom_terms or []
        rules_by_sheet_column: dict[tuple[str, str], ColumnRule] = {
            (rule.sheet, rule.column.upper()): rule for rule in column_rules
        }
        candidates: list[Candidate] = []
        warnings: list[str] = []
        seen: set[tuple[str, str, int, int, SensitiveType]] = set()

        workbook = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=True, rich_text=True)
        try:
            for sheet_name in sheet_names:
                worksheet = workbook[sheet_name]
                for row in worksheet.iter_rows():
                    for cell in row:
                        if cell.value is None or cell.data_type == "f":
                            continue
                        rule = rules_by_sheet_column.get((sheet_name, get_column_letter(cell.column)))
                        if restrict_to_column_rules and rule is None:
                            continue
                        if rule and cell.row <= rule.header_row:
                            continue
                        if not isinstance(cell.value, str):
                            if rule and rule.sensitive_type == SensitiveType.FINANCIAL_AMOUNT and "%" in (cell.number_format or ""):
                                warnings.append(
                                    f"{sheet_name}!{cell.coordinate} 使用百分比格式，已跳过，避免将比例误当金额。"
                                )
                                continue
                            if (
                                rule
                                and rule.sensitive_type == SensitiveType.FINANCIAL_AMOUNT
                                and rule.sensitive_type in enabled_types
                                and isinstance(cell.value, (int, float, Decimal))
                                and not isinstance(cell.value, bool)
                                and not isinstance(cell.value, (date, datetime, time))
                                and not cell.is_date
                            ):
                                self._append_typed_candidate(
                                    candidates,
                                    seen,
                                    sheet_name,
                                    cell.coordinate,
                                    cell.value,
                                    SensitiveType.FINANCIAL_AMOUNT,
                                    0.95 if rule.source == "用户指定列" else 0.82,
                                    rule.source,
                                )
                                continue
                            if rule and rule.sensitive_type == SensitiveType.BANK_ACCOUNT:
                                warnings.append(
                                    f"{sheet_name}!{cell.coordinate} 疑似数值型银行账号，默认未处理；请将其转为文本后重试。"
                                )
                            continue
                        text = str(cell.value)
                        if not text:
                            continue

                        if rule and rule.sensitive_type in WHOLE_CELL_COLUMN_TYPES and rule.sensitive_type in enabled_types:
                            if rule.sensitive_type == SensitiveType.BANK_ACCOUNT:
                                digits = re.sub(r"[ -]", "", text)
                                if not digits.isdigit() or not 12 <= len(digits) <= 19:
                                    warnings.append(
                                        f"{sheet_name}!{cell.coordinate} 位于银行账号列，但不符合 12–19 位文本标识符规则，已跳过。"
                                    )
                                    continue
                            self._append_candidate(
                                candidates,
                                seen,
                                sheet_name,
                                cell.coordinate,
                                text,
                                rule.sensitive_type,
                                0,
                                len(text),
                                0.95 if rule.source == "用户指定列" else 0.82,
                                rule.source,
                            )
                            continue

                        if rule and rule.sensitive_type == SensitiveType.FINANCIAL_AMOUNT and rule.sensitive_type in enabled_types:
                            for match in re.finditer(r"(?<![\w.])[-+]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?![\w.])", text):
                                self._append_candidate(
                                    candidates, seen, sheet_name, cell.coordinate, text,
                                    SensitiveType.FINANCIAL_AMOUNT, match.start(), match.end(),
                                    0.85, f"{rule.source}；文本中的金额数字片段",
                                )
                            continue

                        allow_bank = bool(rule and rule.sensitive_type == SensitiveType.BANK_ACCOUNT)
                        for match in detect_regex(text, enabled_types, allow_bank_account=allow_bank):
                            self._append_candidate(
                                candidates,
                                seen,
                                sheet_name,
                                cell.coordinate,
                                text,
                                match.sensitive_type,
                                match.start,
                                match.end,
                                match.confidence,
                                match.basis,
                            )

                        if SensitiveType.CUSTOM in enabled_types:
                            for term in custom_terms:
                                for start, end in find_custom_term(text, term):
                                    candidate = self._append_candidate(
                                        candidates,
                                        seen,
                                        sheet_name,
                                        cell.coordinate,
                                        text,
                                        SensitiveType.CUSTOM,
                                        start,
                                        end,
                                        1.0,
                                        "用户自定义敏感词",
                                    )
                                    if candidate is not None:
                                        candidate.replacement_override = term.replacement or "【已脱敏】"
        finally:
            workbook.close()

        return ScanResult(candidates=candidates, warnings=list(dict.fromkeys(warnings)))

    @staticmethod
    def _append_candidate(
        candidates: list[Candidate],
        seen: set[tuple[str, str, int, int, SensitiveType]],
        sheet: str,
        coordinate: str,
        cell_value: str,
        sensitive_type: SensitiveType,
        start: int,
        end: int,
        confidence: float,
        basis: str,
    ) -> Candidate | None:
        key = (sheet, coordinate, start, end, sensitive_type)
        if key in seen:
            return None
        seen.add(key)
        fragment = cell_value[start:end]
        digest = hashlib.sha256(
            f"{sheet}\0{coordinate}\0{start}\0{end}\0{sensitive_type.value}\0{fragment}".encode("utf-8")
        ).hexdigest()[:20]
        candidate = Candidate(
            candidate_id=digest,
            sheet=sheet,
            coordinate=coordinate,
            original_value=cell_value,
            original_kind="text",
            safe_preview=safe_preview(fragment, sensitive_type),
            sensitive_type=sensitive_type,
            suggested_method=DEFAULT_METHODS[sensitive_type],
            confidence=confidence,
            basis=basis,
            start=start,
            end=end,
        )
        candidates.append(candidate)
        return candidate

    @staticmethod
    def _append_typed_candidate(
        candidates: list[Candidate],
        seen: set[tuple[str, str, int, int, SensitiveType]],
        sheet: str,
        coordinate: str,
        cell_value: object,
        sensitive_type: SensitiveType,
        confidence: float,
        basis: str,
    ) -> Candidate | None:
        key = (sheet, coordinate, 0, 0, sensitive_type)
        if key in seen:
            return None
        seen.add(key)
        kind = value_kind(cell_value)
        digest = hashlib.sha256(
            f"{sheet}\0{coordinate}\0{sensitive_type.value}\0{kind}\0{cell_value!r}".encode("utf-8")
        ).hexdigest()[:20]
        candidate = Candidate(
            candidate_id=digest,
            sheet=sheet,
            coordinate=coordinate,
            original_value=cell_value,
            original_kind=kind,
            safe_preview=safe_preview(cell_value, sensitive_type),
            sensitive_type=sensitive_type,
            suggested_method=DEFAULT_METHODS[sensitive_type],
            confidence=confidence,
            basis=basis,
            start=0,
            end=0,
            replacement_value=888,
            replacement_kind="integer",
        )
        candidates.append(candidate)
        return candidate


def group_candidate_counts(candidates: list[Candidate]) -> dict[SensitiveType, int]:
    counts: dict[SensitiveType, int] = defaultdict(int)
    for candidate in candidates:
        if candidate.enabled:
            counts[candidate.sensitive_type] += 1
    return dict(counts)


def make_manual_candidate(data: bytes, sheet: str, coordinate: str, sensitive_type: SensitiveType) -> Candidate:
    workbook = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=True, rich_text=True)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError("工作表不存在。")
        cell = workbook[sheet][coordinate.upper().strip()]
        if cell.data_type == "f":
            raise ValueError("不能把公式单元格加入处理计划。")
        if sensitive_type == SensitiveType.FINANCIAL_AMOUNT:
            if (
                not isinstance(cell.value, (int, float, Decimal))
                or isinstance(cell.value, bool)
                or isinstance(cell.value, (date, datetime, time))
                or cell.is_date
            ):
                raise ValueError("财务金额必须是非日期、非布尔的数值单元格。")
            candidates: list[Candidate] = []
            DetectionCoordinator._append_typed_candidate(
                candidates, set(), sheet, cell.coordinate, cell.value,
                sensitive_type, 1.0, "用户手动补充",
            )
            return candidates[0]
        if not isinstance(cell.value, str) or not cell.value:
            raise ValueError("只能手动补充非空文本，或财务金额数值单元格。")
        text = str(cell.value)
        candidates: list[Candidate] = []
        DetectionCoordinator._append_candidate(
            candidates,
            set(),
            sheet,
            cell.coordinate,
            text,
            sensitive_type,
            0,
            len(text),
            1.0,
            "用户手动补充",
        )
        return candidates[0]
    finally:
        workbook.close()
