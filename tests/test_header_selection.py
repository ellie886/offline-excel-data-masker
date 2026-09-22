from __future__ import annotations

import unittest
from io import BytesIO

from openpyxl import Workbook

from offline_masker.detectors.coordinator import DetectionCoordinator
from offline_masker.domain.enums import SensitiveType
from offline_masker.domain.models import ColumnRule
from offline_masker.excel.inspector import WorkbookInspector


def workbook_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


class HeaderSelectionTests(unittest.TestCase):
    def test_detects_header_row_independently_for_each_sheet(self) -> None:
        workbook = Workbook()
        first = workbook.active
        first.title = "客户"
        first.append(["客户明细"])
        first.append([])
        first.append(["客户名称", "手机号", "邮箱", "金额"])
        first.append(["测试甲公司", "13812345678", "a@example.com", 100])

        second = workbook.create_sheet("供应商")
        second.append([])
        second.append(["供应商名称", "联系人", "联系电话"])
        second.append(["测试乙公司", "张三", "13912345678"])

        detected = WorkbookInspector.detect_header_rows(workbook_bytes(workbook), ["客户", "供应商"])

        self.assertEqual(detected["客户"].row, 3)
        self.assertEqual(detected["供应商"].row, 2)
        self.assertGreater(detected["客户"].confidence, 0.6)

    def test_restricted_scan_only_reads_selected_columns(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "明细"
        worksheet.append(["说明"])
        worksheet.append(["手机号", "备注"])
        worksheet.append(["13812345678", "未选择列中的邮箱 hidden@example.com"])
        data = workbook_bytes(workbook)

        result = DetectionCoordinator().scan(
            data,
            ["明细"],
            {SensitiveType.PHONE, SensitiveType.EMAIL},
            [
                ColumnRule(
                    sheet="明细",
                    column="A",
                    sensitive_type=SensitiveType.PHONE,
                    header_row=2,
                    source="用户指定列",
                )
            ],
            restrict_to_column_rules=True,
        )

        self.assertEqual(len(result.candidates), 1)
        self.assertEqual(result.candidates[0].coordinate, "A3")
        self.assertEqual(result.candidates[0].sensitive_type, SensitiveType.PHONE)


if __name__ == "__main__":
    unittest.main()
