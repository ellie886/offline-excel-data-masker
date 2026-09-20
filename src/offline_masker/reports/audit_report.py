from __future__ import annotations

from collections import Counter
from copy import copy
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from offline_masker.domain.models import Candidate, ValidationResult
from offline_masker.security.hashing import short_fingerprint


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
ERROR_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN_GRAY = Side(style="thin", color="D9E1F2")


def build_audit_report(
    *,
    processed_at: datetime,
    source_filename: str,
    output_filename: str,
    scanned_sheets: list[str],
    candidates: list[Candidate],
    validation: ValidationResult,
    warnings: list[str],
) -> bytes:
    selected = [candidate for candidate in candidates if candidate.enabled]
    workbook = Workbook()
    summary = workbook.active
    summary.title = "摘要"
    details = workbook.create_sheet("处理明细")
    checks = workbook.create_sheet("校验结果")

    summary.sheet_view.showGridLines = False
    details.sheet_view.showGridLines = False
    checks.sheet_view.showGridLines = False

    summary.append(["脱敏处理及校验报告"])
    summary["A1"].font = Font(name="Arial", size=14, bold=True, color="1F1F1F")
    summary.append([])
    rows = [
        ("处理时间", processed_at.isoformat(timespec="seconds")),
        ("原文件名称", source_filename),
        ("输出文件名称", output_filename),
        ("扫描工作表", "、".join(scanned_sheets)),
        ("实际替换数量", len(selected)),
        ("最终处理状态", "成功" if validation.passed else "校验失败"),
    ]
    for label, value in rows:
        summary.append([label, value])
    for row in range(3, 3 + len(rows)):
        summary.cell(row, 1).font = Font(name="Arial", bold=True)
        summary.cell(row, 1).fill = SECTION_FILL

    summary.append([])
    summary.append(["各类型替换数量", "数量"])
    _style_header(summary, summary.max_row, 2)
    counts = Counter(candidate.sensitive_type.value for candidate in selected)
    for sensitive_type, count in sorted(counts.items()):
        summary.append([sensitive_type, count])

    summary.append([])
    summary.append(["风险提示及未处理项目"])
    summary[summary.max_row][0].fill = WARNING_FILL
    summary[summary.max_row][0].font = Font(name="Arial", bold=True)
    if warnings:
        for warning in warnings:
            summary.append([warning])
    else:
        summary.append(["无额外风险提示。"])

    details.append(["工作表", "单元格", "安全预览", "类型", "脱敏方式", "识别依据", "原值指纹"])
    _style_header(details, 1, 7)
    for candidate in selected:
        fragment = candidate.original_value[candidate.start : candidate.end]
        details.append(
            [
                candidate.sheet,
                candidate.coordinate,
                candidate.safe_preview,
                candidate.sensitive_type.value,
                candidate.effective_method.value,
                candidate.basis,
                short_fingerprint(fragment),
            ]
        )
    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions

    checks.append(["校验项目", "结果", "说明"])
    _style_header(checks, 1, 3)
    for check in validation.checks:
        checks.append([check.name, "通过" if check.passed else "失败", check.detail])
        if not check.passed:
            for cell in checks[checks.max_row]:
                cell.fill = ERROR_FILL

    for worksheet, widths in (
        (summary, {"A": 28, "B": 55}),
        (details, {"A": 22, "B": 14, "C": 24, "D": 18, "E": 18, "F": 46, "G": 18}),
        (checks, {"A": 34, "B": 12, "C": 68}),
    ):
        for column, width in widths.items():
            worksheet.column_dimensions[column].width = width
        for row in worksheet.iter_rows():
            for cell in row:
                font = copy(cell.font)
                font.name = "Arial"
                font.sz = 10
                cell.font = font
                cell.alignment = Alignment(vertical="center", wrap_text=False)
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _style_header(worksheet, row: int, columns: int) -> None:
    for column in range(1, columns + 1):
        cell = worksheet.cell(row, column)
        cell.fill = HEADER_FILL
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=THIN_GRAY)
