from __future__ import annotations

import unittest
import xml.etree.ElementTree as ET
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from offline_masker.detectors.coordinator import DetectionCoordinator
from offline_masker.domain.enums import SensitiveType
from offline_masker.domain.models import ColumnRule
from offline_masker.excel.inspector import WorkbookInspector
from offline_masker.services.processing_service import ProcessingService


def build_workbook() -> bytes:
    workbook = Workbook()
    customers = workbook.active
    customers.title = "客户明细"
    customers.append(["客户名称", "手机号", "金额", "计算结果", "邮箱", "说明", ""])
    customers.append(["北京甲公司", "13812345678", 100, "=SUM(C2:C3)", "test.user@example.com", "重点客户", None])
    customers.append(["上海乙公司", "13987654321", 200, None, None, None, None])
    customers["A2"].font = Font(name="Arial", bold=True, color="FFFFFF")
    customers["A2"].fill = PatternFill("solid", fgColor="1F4E78")
    customers.column_dimensions["A"].width = 24
    customers.row_dimensions[3].hidden = True
    customers.freeze_panes = "A2"
    customers.auto_filter.ref = "A1:F3"
    customers.merge_cells("F1:G1")

    hidden = workbook.create_sheet("隐藏附表")
    hidden.sheet_state = "hidden"
    hidden.append(["客户名称", "合同编号"])
    hidden.append(["北京甲公司", "HT-2026-001"])

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def convert_a2_to_shared_string(source: bytes) -> bytes:
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ET.register_namespace("", main_ns)

    with ZipFile(BytesIO(source)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    sheet_root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    cell = sheet_root.find(f".//{{{main_ns}}}c[@r='A2']")
    assert cell is not None
    inline = cell.find(f"{{{main_ns}}}is")
    assert inline is not None
    value = "".join(node.text or "" for node in inline.iter(f"{{{main_ns}}}t"))
    cell.remove(inline)
    cell.attrib["t"] = "s"
    ET.SubElement(cell, f"{{{main_ns}}}v").text = "0"
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(sheet_root, encoding="utf-8", xml_declaration=True)

    strings = ET.Element(f"{{{main_ns}}}sst", {"count": "1", "uniqueCount": "1"})
    item = ET.SubElement(strings, f"{{{main_ns}}}si")
    ET.SubElement(item, f"{{{main_ns}}}t").text = value
    parts["xl/sharedStrings.xml"] = ET.tostring(strings, encoding="utf-8", xml_declaration=True)

    content_root = ET.fromstring(parts["[Content_Types].xml"])
    ET.SubElement(
        content_root,
        f"{{{content_ns}}}Override",
        {
            "PartName": "/xl/sharedStrings.xml",
            "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
        },
    )
    parts["[Content_Types].xml"] = ET.tostring(content_root, encoding="utf-8", xml_declaration=True)

    rels_root = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
    ET.SubElement(
        rels_root,
        f"{{{rel_ns}}}Relationship",
        {
            "Id": "rIdSharedStringsTest",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
            "Target": "sharedStrings.xml",
        },
    )
    parts["xl/_rels/workbook.xml.rels"] = ET.tostring(rels_root, encoding="utf-8", xml_declaration=True)

    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    return output.getvalue()


class EndToEndTests(unittest.TestCase):
    def test_shared_string_cell_is_supported(self) -> None:
        source = convert_a2_to_shared_string(build_workbook())
        scan = DetectionCoordinator().scan(
            source,
            ["客户明细"],
            {SensitiveType.CUSTOMER},
            [ColumnRule("客户明细", "A", SensitiveType.CUSTOMER)],
        )
        self.assertEqual(len(scan.candidates), 2)
        result = ProcessingService().process(
            source=source,
            source_filename="shared.xlsx",
            scanned_sheets=["客户明细"],
            candidates=scan.candidates,
        )
        self.assertTrue(result.validation.passed)
        workbook = load_workbook(BytesIO(result.output_bytes))
        try:
            self.assertEqual(workbook["客户明细"]["A2"].value, "客户001")
        finally:
            workbook.close()

    def test_masking_preserves_structure_formulas_numbers_and_consistency(self) -> None:
        source = build_workbook()
        inspection = WorkbookInspector().inspect(source, "测试工作簿.xlsx")
        self.assertEqual([sheet.name for sheet in inspection.sheets], ["客户明细", "隐藏附表"])
        self.assertTrue(any("隐藏状态" in warning for warning in inspection.warnings))

        rules = [
            ColumnRule("客户明细", "A", SensitiveType.CUSTOMER),
            ColumnRule("隐藏附表", "A", SensitiveType.CUSTOMER),
            ColumnRule("隐藏附表", "B", SensitiveType.CONTRACT),
        ]
        scan = DetectionCoordinator().scan(
            source,
            ["客户明细", "隐藏附表"],
            {
                SensitiveType.CUSTOMER,
                SensitiveType.CONTRACT,
                SensitiveType.PHONE,
                SensitiveType.EMAIL,
            },
            rules,
        )
        self.assertEqual(len(scan.candidates), 7)

        result = ProcessingService().process(
            source=source,
            source_filename="测试工作簿.xlsx",
            scanned_sheets=["客户明细", "隐藏附表"],
            candidates=scan.candidates,
            warnings=inspection.warnings + scan.warnings,
        )
        failed = [check for check in result.validation.checks if not check.passed]
        self.assertEqual(failed, [])
        self.assertEqual(result.replacement_count, 7)

        workbook = load_workbook(BytesIO(result.output_bytes), data_only=False)
        try:
            self.assertEqual(workbook.sheetnames, ["客户明细", "隐藏附表"])
            self.assertEqual(workbook["隐藏附表"].sheet_state, "hidden")
            self.assertEqual(workbook["客户明细"]["A2"].value, "客户001")
            self.assertEqual(workbook["隐藏附表"]["A2"].value, "客户001")
            self.assertEqual(workbook["客户明细"]["B2"].value, "138****5678")
            self.assertEqual(workbook["客户明细"]["E2"].value, "t*******r@example.com")
            self.assertEqual(workbook["客户明细"]["C2"].value, 100)
            self.assertEqual(workbook["客户明细"]["C3"].value, 200)
            self.assertEqual(workbook["客户明细"]["D2"].value, "=SUM(C2:C3)")
            self.assertTrue(workbook["客户明细"].row_dimensions[3].hidden)
            self.assertEqual(workbook["客户明细"].freeze_panes, "A2")
            self.assertIn("F1:G1", {str(item) for item in workbook["客户明细"].merged_cells.ranges})
        finally:
            workbook.close()

        with ZipFile(BytesIO(result.report_bytes)) as archive:
            report_payload = b"".join(archive.read(name) for name in archive.namelist())
        for secret in ("北京甲公司", "上海乙公司", "13812345678", "test.user@example.com"):
            self.assertNotIn(secret.encode("utf-8"), report_payload)


if __name__ == "__main__":
    unittest.main()
