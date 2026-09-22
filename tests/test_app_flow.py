from __future__ import annotations

import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from streamlit.testing.v1 import AppTest


class AppFlowTests(unittest.TestCase):
    def test_multiple_sheets_use_independent_detected_header_rows(self) -> None:
        workbook = Workbook()
        first = workbook.active
        first.title = "客户"
        first.append(["客户明细"])
        first.append([])
        first.append(["客户名称", "手机号", "金额"])
        first.append(["测试甲公司", "13812345678", 100])
        second = workbook.create_sheet("供应商")
        second.append([])
        second.append(["供应商名称", "联系人", "联系电话"])
        second.append(["测试乙公司", "张三", "13912345678"])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        app_path = Path(__file__).resolve().parents[1] / "app.py"
        app = AppTest.from_file(str(app_path)).run(timeout=30)
        app.get("file_uploader")[0].upload("multi-sheet.xlsx", buffer.getvalue()).run(timeout=30)

        header_inputs = [item.value for item in app.number_input if item.label == "表头所在行"]
        self.assertEqual(header_inputs, [3, 2])
        self.assertEqual(len(app.exception), 0)

    def test_upload_scan_confirm_process_and_export(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "客户明细"
        worksheet.append(["客户名称", "手机号", "邮箱", "金额"])
        worksheet.append(["北京甲公司", "13812345678", "test@example.com", 100])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        app_path = Path(__file__).resolve().parents[1] / "app.py"
        app = AppTest.from_file(str(app_path)).run(timeout=30)
        app.get("file_uploader")[0].upload("demo.xlsx", buffer.getvalue()).run(timeout=30)
        next(button for button in app.button if button.label == "扫描敏感信息").click().run(timeout=30)

        confirm = next(item for item in app.checkbox if item.label.startswith("我已确认只修改"))
        confirm.check().run(timeout=30)
        next(button for button in app.button if button.label == "执行脱敏并校验").click().run(timeout=30)

        self.assertEqual(len(app.exception), 0)
        self.assertEqual([item.value for item in app.success], ["处理和关键校验均已通过，共替换 4 项。"])
        self.assertEqual(
            [item.label for item in app.get("download_button")],
            ["下载脱敏文件", "下载脱敏报告"],
        )


if __name__ == "__main__":
    unittest.main()
