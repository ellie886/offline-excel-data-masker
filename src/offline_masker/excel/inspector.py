from __future__ import annotations

import re
import xml.etree.ElementTree as ET
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from offline_masker.detectors.header_detector import infer_header_type
from offline_masker.domain.exceptions import UnsafeWorkbookError
from offline_masker.domain.models import HeaderDetection, SheetInfo, WorkbookInspection
from offline_masker.security.file_guard import FileLimits, validate_xlsx_package
from offline_masker.excel.ooxml_patcher import MAIN_NS, resolve_sheet_paths


RISK_PARTS: tuple[tuple[str, str], ...] = (
    ("xl/charts/", "图表数据缓存"),
    ("xl/drawings/", "文本框、形状或绘图内容"),
    ("xl/pivot", "透视表或透视表缓存"),
    ("xl/threadedComments/", "线程批注"),
    ("xl/comments", "普通批注"),
    ("xl/embeddings/", "嵌入对象或附件"),
    ("customXml/", "自定义 XML 内容"),
    ("xl/connections.xml", "外部连接或 Power Query"),
    ("xl/externalLinks/", "外部工作簿连接"),
    ("xl/queryTables/", "Power Query 或查询表"),
    ("xl/queries/", "Power Query 查询"),
    ("xl/model/", "Excel 数据模型"),
)

GENERIC_HEADER_PATTERN = re.compile(
    r"名称|编号|代码|日期|时间|期间|年度|月份|金额|收入|成本|数量|单价|税率|币种|汇率|状态|备注|描述|地区|国家|科目|类型|属性|订单|发票|合同|物料"
)


class WorkbookInspector:
    def __init__(self, *, max_scannable_cells: int = 1_000_000, limits: FileLimits | None = None) -> None:
        self.max_scannable_cells = max_scannable_cells
        self.limits = limits or FileLimits()

    def inspect(self, data: bytes, filename: str) -> WorkbookInspection:
        warnings = validate_xlsx_package(data, filename, self.limits)
        try:
            with ZipFile(BytesIO(data)) as archive:
                package_parts = sorted(archive.namelist())
                sheet_paths = resolve_sheet_paths(archive)
                formula_cache_counts: dict[str, int] = {}
                protection_risks: list[str] = []
                workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
                workbook_protection = workbook_root.find(f"{{{MAIN_NS}}}workbookProtection")
                if workbook_protection is not None and bool(workbook_protection.attrib):
                    protection_risks.append("工作簿结构或修订保护/密码保护")
                for sheet_name, path in sheet_paths.items():
                    root = ET.fromstring(archive.read(path))
                    formula_cache_counts[sheet_name] = sum(
                        1 for cell in root.findall(f".//{{{MAIN_NS}}}c")
                        if cell.find(f"{{{MAIN_NS}}}f") is not None and cell.find(f"{{{MAIN_NS}}}v") is not None
                    )
                    if root.find(f"{{{MAIN_NS}}}sheetProtection") is not None:
                        protection_risks.append(f"工作表“{sheet_name}”存在密码或编辑保护")
        except (ET.ParseError, KeyError, UnsafeWorkbookError) as exc:
            raise UnsafeWorkbookError("工作簿包含无法安全解析的数据承载组件。") from exc
        unsupported = [message for prefix, message in RISK_PARTS if any(name.startswith(prefix) for name in package_parts)]
        unsupported.extend(protection_risks)

        try:
            workbook = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=True, rich_text=True)
        except Exception as exc:
            raise UnsafeWorkbookError(f"无法安全读取工作簿：{type(exc).__name__}。") from exc

        sheets: list[SheetInfo] = []
        total_cells = 0
        try:
            for worksheet in workbook.worksheets:
                estimated_cells = worksheet.max_row * worksheet.max_column
                total_cells += estimated_cells
                if total_cells > self.max_scannable_cells:
                    raise UnsafeWorkbookError(
                        f"工作簿使用区域约含 {total_cells:,} 个单元格，超过第一版 {self.max_scannable_cells:,} 个的扫描限制。"
                    )

                formula_count = 0
                comment_count = 0
                for row in worksheet.iter_rows():
                    for cell in row:
                        if cell.data_type == "f":
                            formula_count += 1
                        if cell.comment is not None:
                            comment_count += 1

                hidden_rows = sum(1 for dimension in worksheet.row_dimensions.values() if dimension.hidden)
                hidden_columns = 0
                for key, dimension in worksheet.column_dimensions.items():
                    if dimension.hidden:
                        start = dimension.min or worksheet[key][0].column
                        end = dimension.max or start
                        hidden_columns += end - start + 1

                if worksheet.sheet_state != "visible":
                    warnings.append(f"工作表“{worksheet.title}”处于隐藏状态（{worksheet.sheet_state}）。")
                if hidden_rows:
                    warnings.append(f"工作表“{worksheet.title}”包含 {hidden_rows} 个隐藏行定义。")
                if hidden_columns:
                    warnings.append(f"工作表“{worksheet.title}”包含 {hidden_columns} 个隐藏列定义。")
                if comment_count:
                    warnings.append(f"工作表“{worksheet.title}”包含 {comment_count} 条普通批注，第一版只提示、不扫描批注正文。")

                sheets.append(
                    SheetInfo(
                        name=worksheet.title,
                        state=worksheet.sheet_state,
                        max_row=worksheet.max_row,
                        max_column=worksheet.max_column,
                        formula_count=formula_count,
                        formula_cache_count=formula_cache_counts.get(worksheet.title, 0),
                        merged_range_count=len(worksheet.merged_cells.ranges),
                        hidden_row_count=hidden_rows,
                        hidden_column_count=hidden_columns,
                        comment_count=comment_count,
                    )
                )
        finally:
            workbook.close()

        if unsupported:
            warnings.extend(f"未处理内容：{item}。" for item in unsupported)
        return WorkbookInspection(
            filename=filename,
            size_bytes=len(data),
            sheets=sheets,
            warnings=_deduplicate(warnings),
            unsupported_items=unsupported,
            package_parts=package_parts,
            blocking_risks=_deduplicate(unsupported),
        )

    @staticmethod
    def list_headers(data: bytes, sheet_name: str, header_row: int = 1) -> list[tuple[str, str]]:
        workbook = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=True, rich_text=True)
        try:
            worksheet = workbook[sheet_name]
            headers: list[tuple[str, str]] = []
            for column_index in range(1, worksheet.max_column + 1):
                value = worksheet.cell(header_row, column_index).value
                headers.append((get_column_letter(column_index), "" if value is None else str(value)))
            return headers
        finally:
            workbook.close()

    @staticmethod
    def detect_header_rows(
        data: bytes,
        sheet_names: list[str],
        *,
        max_scan_rows: int = 50,
    ) -> dict[str, HeaderDetection]:
        workbook = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=True, rich_text=True)
        detections: dict[str, HeaderDetection] = {}
        try:
            for sheet_name in sheet_names:
                worksheet = workbook[sheet_name]
                candidates: list[tuple[float, int, int, int]] = []
                upper_row = min(max_scan_rows, worksheet.max_row)
                upper_column = min(worksheet.max_column, 300)
                for row_index in range(1, upper_row + 1):
                    values = [worksheet.cell(row_index, column).value for column in range(1, upper_column + 1)]
                    nonempty = [value for value in values if value is not None and str(value).strip()]
                    if not nonempty:
                        continue
                    text_values = [str(value).strip() for value in nonempty if isinstance(value, str)]
                    sensitive_hits = sum(1 for value in nonempty if infer_header_type(value) is not None)
                    generic_hits = sum(1 for value in nonempty if GENERIC_HEADER_PATTERN.search(str(value)))
                    unique_ratio = len(set(str(value).strip() for value in nonempty)) / len(nonempty)
                    numeric_ratio = sum(1 for value in nonempty if isinstance(value, (int, float))) / len(nonempty)
                    next_nonempty = 0
                    if row_index < worksheet.max_row:
                        next_nonempty = sum(
                            1
                            for column in range(1, upper_column + 1)
                            if worksheet.cell(row_index + 1, column).value is not None
                            and str(worksheet.cell(row_index + 1, column).value).strip()
                        )
                    score = (
                        len(nonempty) * 2.0
                        + len(text_values) * 1.5
                        + sensitive_hits * 5.0
                        + generic_hits * 2.5
                        + unique_ratio * 3.0
                        + min(next_nonempty, len(nonempty)) * 0.35
                        - numeric_ratio * 4.0
                    )
                    if len(nonempty) == 1:
                        score -= 8.0
                    candidates.append((score, row_index, len(nonempty), sensitive_hits + generic_hits))

                if not candidates:
                    detections[sheet_name] = HeaderDetection(
                        sheet=sheet_name,
                        row=1,
                        confidence=0.0,
                        basis="未发现非空行，默认使用第 1 行",
                    )
                    continue

                candidates.sort(reverse=True)
                best_score, best_row, nonempty_count, keyword_hits = candidates[0]
                second_score = candidates[1][0] if len(candidates) > 1 else 0.0
                margin = max(0.0, best_score - second_score)
                confidence = min(0.99, 0.62 + min(0.25, margin / max(abs(best_score), 1.0)) + min(0.12, keyword_hits * 0.02))
                detections[sheet_name] = HeaderDetection(
                    sheet=sheet_name,
                    row=best_row,
                    confidence=confidence,
                    basis=f"第 {best_row} 行包含 {nonempty_count} 个非空单元格，命中 {keyword_hits} 个表头关键词",
                )
        finally:
            workbook.close()
        return detections


def _deduplicate(items: list[str]) -> list[str]:
    return list(dict.fromkeys(items))
