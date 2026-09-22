from __future__ import annotations

import hashlib
import xml.etree.ElementTree as ET
from datetime import date, datetime, time
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

from offline_masker.domain.models import PatchOperation, ValidationCheck, ValidationResult
from offline_masker.security.hashing import value_kind
from offline_masker.excel.ooxml_patcher import MAIN_NS, parse_xml_preserving_namespaces, resolve_sheet_paths


class ExcelValidator:
    """Validate an explicit allowlist: targets, shared-string cleanup and formula cache metadata."""

    def validate(
        self,
        source: bytes,
        output: bytes,
        operations: list[PatchOperation],
        modified_parts: set[str],
        actual_replacement_count: int,
        *,
        source_sha256_before: str | None = None,
    ) -> ValidationResult:
        checks: list[ValidationCheck] = []
        source_digest = hashlib.sha256(source).hexdigest()
        checks.append(ValidationCheck(
            "原始文件 SHA-256",
            source_sha256_before in (None, source_digest),
            "处理前后原始文件字节完全一致。" if source_sha256_before in (None, source_digest) else "原始文件哈希发生变化。",
        ))
        try:
            before = self._snapshot(source)
            after = self._snapshot(output)
        except Exception as exc:
            checks.append(ValidationCheck("输出文件可读取", False, f"无法重新读取输出文件：{type(exc).__name__}。"))
            return ValidationResult(checks)
        checks.append(ValidationCheck("输出文件可读取", True, "openpyxl 已成功重新打开输出文件。"))

        for name, key in (
            ("工作表名称和顺序", "sheet_names"),
            ("工作表隐藏状态", "sheet_states"),
            ("工作表行列范围", "dimensions"),
            ("公式文本逐单元格", "formulas"),
            ("合并区域", "merged"),
            ("行高、列宽、隐藏与分组", "dimensions_meta"),
            ("冻结窗格、筛选与打印设置", "sheet_features"),
            ("条件格式", "conditional_formatting"),
            ("数据验证", "data_validations"),
            ("表格定义", "tables"),
            ("超链接", "hyperlinks"),
            ("工作簿内部组件清单", "package_parts"),
        ):
            checks.append(self._equal_check(name, before[key], after[key]))

        targets = {(item.sheet, item.coordinate): item for item in operations}
        non_target_differences: list[str] = []
        all_refs = set(before["cells"]) | set(after["cells"])
        for reference in sorted(all_refs):
            sheet, coordinate = reference
            if (sheet, coordinate) in targets:
                continue
            if before["cells"].get(reference) != after["cells"].get(reference):
                non_target_differences.append(f"{sheet}!{coordinate}")
        checks.append(self._locations_check("非目标单元格值与类型", non_target_differences))

        style_differences = [
            f"{sheet}!{coordinate}" for sheet, coordinate in sorted(set(before["styles"]) | set(after["styles"]))
            if before["styles"].get((sheet, coordinate)) != after["styles"].get((sheet, coordinate))
        ]
        checks.append(self._locations_check("样式 ID、数字格式、字体、填充、边框、对齐与保护", style_differences))

        target_differences: list[str] = []
        for key, operation in targets.items():
            actual = after["cells"].get(key)
            if actual is None or actual[0] != operation.replacement_kind or actual[1] != operation.replacement_value:
                target_differences.append(f"{operation.sheet}!{operation.coordinate}")
        checks.append(self._locations_check("目标单元格替换结果及类型", target_differences))

        financial_errors: list[str] = []
        for operation in operations:
            if operation.replacement_kind == "integer" and operation.replacement_value == 888:
                actual = after["cells"].get((operation.sheet, operation.coordinate))
                if actual is None or actual[0] != "integer" or actual[1] != 888:
                    financial_errors.append(f"{operation.sheet}!{operation.coordinate}")
        checks.append(self._locations_check("财务金额为数值型 888（非文本）", financial_errors))

        expected_count = sum(item.candidate_count for item in operations)
        checks.append(ValidationCheck("实际替换数量", actual_replacement_count == expected_count,
                                      f"确认 {expected_count} 项，实际执行 {actual_replacement_count} 项。"))
        checks.append(self._validate_untouched_parts(source, output, modified_parts))
        checks.append(self._validate_changed_worksheet_xml(source, output, operations, modified_parts))
        checks.append(self._validate_workbook_xml_allowlist(source, output))
        checks.append(self._validate_shared_strings_allowlist(source, output, operations))
        checks.append(self._validate_no_sensitive_residue(output, operations))
        return ValidationResult(checks)

    @staticmethod
    def _snapshot(data: bytes) -> dict[str, object]:
        workbook = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=True, rich_text=True)
        result: dict[str, object] = {
            "sheet_names": tuple(workbook.sheetnames), "sheet_states": {}, "dimensions": {},
            "formulas": {}, "cells": {}, "styles": {}, "merged": {}, "dimensions_meta": {},
            "sheet_features": {}, "conditional_formatting": {}, "data_validations": {},
            "tables": {}, "hyperlinks": {},
        }
        try:
            for ws in workbook.worksheets:
                result["sheet_states"][ws.title] = ws.sheet_state
                result["dimensions"][ws.title] = (ws.max_row, ws.max_column)
                for row in ws.iter_rows():
                    for cell in row:
                        key = (ws.title, cell.coordinate)
                        value = cell.value
                        if cell.data_type == "f":
                            result["formulas"][key] = str(value)
                        if value is not None:
                            result["cells"][key] = (value_kind(value), value, cell.data_type)
                        if value is not None or cell.has_style:
                            result["styles"][key] = (
                                cell.style_id, tuple(cell._style) if cell._style is not None else (),
                                cell.number_format, str(cell.font), str(cell.fill), str(cell.border),
                                str(cell.alignment), str(cell.protection),
                            )
                        if cell.hyperlink:
                            link = cell.hyperlink
                            result["hyperlinks"][key] = (link.target, link.location, link.display, link.tooltip)
                result["merged"][ws.title] = tuple(sorted(str(item) for item in ws.merged_cells.ranges))
                result["dimensions_meta"][ws.title] = (
                    tuple(sorted((i, d.hidden, d.height, d.outlineLevel, d.collapsed) for i, d in ws.row_dimensions.items())),
                    tuple(sorted((i, d.hidden, d.width, d.min, d.max, d.outlineLevel, d.collapsed) for i, d in ws.column_dimensions.items())),
                )
                result["sheet_features"][ws.title] = (
                    str(ws.freeze_panes or ""), ws.auto_filter.ref or "", str(ws.print_area or ""),
                    str(ws.print_title_rows or ""), str(ws.print_title_cols or ""),
                    str(ws.page_margins), str(ws.page_setup), str(ws.print_options), str(ws.sheet_properties.pageSetUpPr),
                )
                result["conditional_formatting"][ws.title] = tuple(
                    (str(key.sqref), tuple(ET.tostring(rule.to_tree(), encoding="unicode") for rule in rules))
                    for key, rules in ws.conditional_formatting._cf_rules.items()
                )
                result["data_validations"][ws.title] = tuple(
                    ET.tostring(item.to_tree(), encoding="unicode") for item in ws.data_validations.dataValidation
                )
                result["tables"][ws.title] = tuple(
                    sorted(ET.tostring(table.to_tree(), encoding="unicode") for table in ws.tables.values())
                )
        finally:
            workbook.close()
        with ZipFile(BytesIO(data)) as archive:
            result["package_parts"] = tuple(sorted(archive.namelist()))
        return result

    @staticmethod
    def _equal_check(name: str, before: object, after: object) -> ValidationCheck:
        if before == after:
            return ValidationCheck(name, True, "处理前后一致。")
        if isinstance(before, dict) and isinstance(after, dict):
            keys = set(before) | set(after)
            changed = [str(key) for key in keys if before.get(key) != after.get(key)]
            return ValidationCheck(name, False, f"发现 {len(changed)} 处差异：{'、'.join(changed[:5])}。")
        return ValidationCheck(name, False, "处理前后不一致。")

    @staticmethod
    def _locations_check(name: str, locations: list[str]) -> ValidationCheck:
        if not locations:
            return ValidationCheck(name, True, "未发现超出白名单的变化。")
        return ValidationCheck(name, False, f"发现 {len(locations)} 处差异：{'、'.join(locations[:8])}。")

    @staticmethod
    def _validate_untouched_parts(source: bytes, output: bytes, modified_parts: set[str]) -> ValidationCheck:
        with ZipFile(BytesIO(source)) as before, ZipFile(BytesIO(output)) as after:
            names = set(before.namelist())
            if names != set(after.namelist()):
                return ValidationCheck("非白名单 OOXML 组件", False, "内部组件清单发生变化。")
            changed = [name for name in sorted(names - modified_parts) if before.read(name) != after.read(name)]
        if changed:
            return ValidationCheck("非白名单 OOXML 组件", False, f"发现 {len(changed)} 个非白名单组件变化：{'、'.join(changed[:5])}。")
        return ValidationCheck("非白名单 OOXML 组件", True, "所有非白名单组件字节完全一致。")

    @staticmethod
    def _validate_no_sensitive_residue(output: bytes, operations: list[PatchOperation]) -> ValidationCheck:
        text_secrets = {str(value) for item in operations for value in item.sensitive_originals if isinstance(value, str) and value}
        number_secrets = {str(value) for item in operations for value in item.sensitive_originals if isinstance(value, (int, float)) and not isinstance(value, bool)}
        findings: list[str] = []
        with ZipFile(BytesIO(output)) as archive:
            for name in archive.namelist():
                if not name.lower().endswith((".xml", ".rels")):
                    continue
                payload = archive.read(name)
                try:
                    root = ET.fromstring(payload)
                except ET.ParseError:
                    findings.append(name)
                    continue
                for node in root.iter():
                    text = node.text or ""
                    if any(secret in text for secret in text_secrets):
                        findings.append(name)
                        break
                    local = node.tag.rsplit("}", 1)[-1] if isinstance(node.tag, str) else ""
                    if local == "v" and text in number_secrets:
                        findings.append(name)
                        break
        if findings:
            return ValidationCheck("OOXML 敏感原值残留扫描", False, f"敏感原值仍可能存在于：{'、'.join(sorted(set(findings))[:8])}。")
        return ValidationCheck("OOXML 敏感原值残留扫描", True, "解压后的 XML/RELS 中未发现已确认的敏感原值。")

    @staticmethod
    def _validate_changed_worksheet_xml(
        source: bytes, output: bytes, operations: list[PatchOperation], modified_parts: set[str]
    ) -> ValidationCheck:
        targets: dict[str, set[str]] = {}
        for item in operations:
            targets.setdefault(item.sheet, set()).add(item.coordinate)
        differences: list[str] = []
        with ZipFile(BytesIO(source)) as before_zip, ZipFile(BytesIO(output)) as after_zip:
            before_paths = resolve_sheet_paths(before_zip)
            after_paths = resolve_sheet_paths(after_zip)
            before_shared = ExcelValidator._shared_values(before_zip)
            after_shared = ExcelValidator._shared_values(after_zip)
            path_to_sheet = {path: sheet for sheet, path in before_paths.items()}
            for path in sorted(name for name in modified_parts if name.startswith("xl/worksheets/")):
                if path not in before_zip.namelist() or path not in after_zip.namelist():
                    differences.append(path)
                    continue
                sheet = path_to_sheet.get(path, path)
                before_root = parse_xml_preserving_namespaces(before_zip.read(path))
                after_root = parse_xml_preserving_namespaces(after_zip.read(path))
                ExcelValidator._normalise_allowed_sheet_changes(before_root, targets.get(sheet, set()), before_shared)
                ExcelValidator._normalise_allowed_sheet_changes(after_root, targets.get(sheet, set()), after_shared)
                if _element_signature(before_root) != _element_signature(after_root):
                    differences.append(sheet)
            if before_paths.keys() != after_paths.keys():
                differences.append("工作表关系")
        if differences:
            return ValidationCheck("已修改工作表 XML 白名单", False, f"目标值、公式缓存和共享索引以外仍有变化：{'、'.join(differences[:6])}。")
        return ValidationCheck("已修改工作表 XML 白名单", True, "仅发现目标值、公式缓存或共享字符串索引变化。")

    @staticmethod
    def _normalise_allowed_sheet_changes(root: ET.Element, targets: set[str], shared: list[str]) -> None:
        for cell in root.findall(f".//{{{MAIN_NS}}}c"):
            if cell.attrib.get("r") in targets:
                cell.attrib["t"] = "__confirmed_target__"
                for tag in (f"{{{MAIN_NS}}}v", f"{{{MAIN_NS}}}is"):
                    node = cell.find(tag)
                    if node is not None:
                        cell.remove(node)
            elif cell.attrib.get("t") == "s":
                node = cell.find(f"{{{MAIN_NS}}}v")
                if node is not None and node.text is not None:
                    try:
                        node.text = "__shared_value__" + shared[int(node.text)]
                    except (ValueError, IndexError):
                        node.text = "__invalid_shared_index__"
            if cell.find(f"{{{MAIN_NS}}}f") is not None:
                cache = cell.find(f"{{{MAIN_NS}}}v")
                if cache is not None:
                    cell.remove(cache)

    @staticmethod
    def _shared_values(archive: ZipFile) -> list[str]:
        if "xl/sharedStrings.xml" not in archive.namelist():
            return []
        root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
        return ["".join(node.text or "" for node in item.iter(f"{{{MAIN_NS}}}t")) for item in root.findall(f"{{{MAIN_NS}}}si")]

    @staticmethod
    def _validate_workbook_xml_allowlist(source: bytes, output: bytes) -> ValidationCheck:
        with ZipFile(BytesIO(source)) as before, ZipFile(BytesIO(output)) as after:
            left = parse_xml_preserving_namespaces(before.read("xl/workbook.xml"))
            right = parse_xml_preserving_namespaces(after.read("xl/workbook.xml"))
        allowed = {"calcMode", "fullCalcOnLoad", "forceFullCalc"}
        left_calc = left.find(f"{{{MAIN_NS}}}calcPr")
        right_calc = right.find(f"{{{MAIN_NS}}}calcPr")
        for node in (left_calc, right_calc):
            if node is not None:
                for attribute in allowed:
                    node.attrib.pop(attribute, None)
        if left_calc is None and right_calc is not None and not right_calc.attrib and not list(right_calc):
            right.remove(right_calc)
        if right_calc is None and left_calc is not None and not left_calc.attrib and not list(left_calc):
            left.remove(left_calc)
        passed = _element_signature(left) == _element_signature(right)
        return ValidationCheck("工作簿计算属性白名单", passed,
                               "仅允许的完整重算属性可能变化。" if passed else "workbook.xml 存在计算属性以外的变化。")

    @staticmethod
    def _validate_shared_strings_allowlist(source: bytes, output: bytes, operations: list[PatchOperation]) -> ValidationCheck:
        with ZipFile(BytesIO(source)) as before, ZipFile(BytesIO(output)) as after:
            before_items = ExcelValidator._shared_item_signatures(before)
            after_items = ExcelValidator._shared_item_signatures(after)
        cursor = 0
        removed: list[str] = []
        for signature, text in before_items:
            if cursor < len(after_items) and signature == after_items[cursor][0]:
                cursor += 1
            else:
                removed.append(text)
        sensitive_text = {str(value) for item in operations for value in item.sensitive_originals if isinstance(value, str) and value}
        passed = cursor == len(after_items) and all(any(secret in text for secret in sensitive_text) for text in removed)
        detail = f"仅删除了 {len(removed)} 个无引用敏感共享字符串。" if passed else "sharedStrings.xml 存在删除敏感无引用项以外的变化。"
        return ValidationCheck("共享字符串白名单", passed, detail)

    @staticmethod
    def _shared_item_signatures(archive: ZipFile) -> list[tuple[tuple[object, ...], str]]:
        if "xl/sharedStrings.xml" not in archive.namelist():
            return []
        root = parse_xml_preserving_namespaces(archive.read("xl/sharedStrings.xml"))
        return [
            (_element_signature(item), "".join(node.text or "" for node in item.iter(f"{{{MAIN_NS}}}t")))
            for item in root.findall(f"{{{MAIN_NS}}}si")
        ]


def report_contains_sensitive_value(report: bytes, operations: list[PatchOperation]) -> list[str]:
    text_secrets = {str(value) for item in operations for value in item.sensitive_originals if isinstance(value, str) and value}
    number_secrets = {str(value) for item in operations for value in item.sensitive_originals if isinstance(value, (int, float)) and not isinstance(value, bool)}
    findings: list[str] = []
    with ZipFile(BytesIO(report)) as archive:
        for name in archive.namelist():
            if not name.lower().endswith((".xml", ".rels")):
                continue
            try:
                root = ET.fromstring(archive.read(name))
            except ET.ParseError:
                findings.append(name)
                continue
            for node in root.iter():
                content = node.text or ""
                if any(secret in content for secret in text_secrets):
                    findings.append(name)
                    break
                local = node.tag.rsplit("}", 1)[-1] if isinstance(node.tag, str) else ""
                if local == "v" and content in number_secrets:
                    findings.append(name)
                    break
    return sorted(set(findings))


def _element_signature(element: ET.Element) -> tuple[object, ...]:
    if not isinstance(element.tag, str):
        return ("#comment", element.text or "")
    text = element.text if element.text and element.text.strip() else ""
    return (element.tag, tuple(sorted(element.attrib.items())), text, tuple(_element_signature(child) for child in list(element)))
