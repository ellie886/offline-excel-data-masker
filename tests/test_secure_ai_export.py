from __future__ import annotations

import hashlib
import unittest
import xml.etree.ElementTree as ET
from datetime import date
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from offline_masker.detectors.coordinator import DetectionCoordinator
from offline_masker.domain.enums import SensitiveType
from offline_masker.domain.exceptions import UnsafeWorkbookError
from offline_masker.domain.models import ColumnRule
from offline_masker.excel.ooxml_patcher import MAIN_NS, OoxmlPatcher
from offline_masker.maskers.replacement_planner import ReplacementPlanner
from offline_masker.security.hashing import value_hash
from offline_masker.services.processing_service import ProcessingService
from offline_masker.validation.excel_validator import ExcelValidator


def save(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def financial_source() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "财务"
    sheet.append(["项目", "金额", "非目标数字", "公式"])
    sheet.append(["收入", 123456.78, 42, "=SUM(B2:B6)"])
    sheet.append(["日期", date(2026, 9, 22), 43, None])
    sheet.append(["比例", 0.15, 44, None])
    sheet.append(["布尔", True, 45, None])
    sheet.append(["空值", None, 46, None])
    sheet["B2"].number_format = '¥#,##0.00'
    sheet["B2"].font = Font(bold=True, color="FFFFFF")
    sheet["B2"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet["B4"].number_format = "0.00%"
    sheet.merge_cells("E1:F1")
    sheet.row_dimensions[5].hidden = True
    sheet.column_dimensions["F"].hidden = True
    validation = DataValidation(type="whole", operator="between", formula1="0", formula2="1000000")
    sheet.add_data_validation(validation)
    validation.add("B2:B6")
    sheet.conditional_formatting.add("B2:B6", CellIsRule(operator="greaterThan", formula=["0"]))
    return save(workbook)


def inject_formula_cache(source: bytes, value: str = "123456.78") -> bytes:
    with ZipFile(BytesIO(source)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
        infos = archive.infolist()
    root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    cell = root.find(f".//{{{MAIN_NS}}}c[@r='D2']")
    assert cell is not None and cell.find(f"{{{MAIN_NS}}}f") is not None
    cached = cell.find(f"{{{MAIN_NS}}}v")
    if cached is None:
        cached = ET.SubElement(cell, f"{{{MAIN_NS}}}v")
    cached.text = value
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    return rebuild(parts, infos)


def rebuild(parts: dict[str, bytes], infos) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
        for info in infos:
            archive.writestr(info, parts[info.filename])
        known = {info.filename for info in infos}
        for name, payload in parts.items():
            if name not in known:
                archive.writestr(name, payload)
    return buffer.getvalue()


def add_part(source: bytes, name: str, payload: bytes = b"<root/>") -> bytes:
    with ZipFile(BytesIO(source)) as archive:
        parts = {item.filename: archive.read(item.filename) for item in archive.infolist()}
        infos = archive.infolist()
    parts[name] = payload
    return rebuild(parts, infos)


def to_shared_string(source: bytes, coordinates: list[str]) -> bytes:
    content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with ZipFile(BytesIO(source)) as archive:
        parts = {item.filename: archive.read(item.filename) for item in archive.infolist()}
        infos = archive.infolist()
    root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    original = None
    for coordinate in coordinates:
        cell = root.find(f".//{{{MAIN_NS}}}c[@r='{coordinate}']")
        assert cell is not None
        inline = cell.find(f"{{{MAIN_NS}}}is")
        assert inline is not None
        text = "".join(item.text or "" for item in inline.iter(f"{{{MAIN_NS}}}t"))
        original = original or text
        assert text == original
        cell.remove(inline)
        cell.set("t", "s")
        ET.SubElement(cell, f"{{{MAIN_NS}}}v").text = "0"
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    strings = ET.Element(f"{{{MAIN_NS}}}sst", {"count": str(len(coordinates)), "uniqueCount": "1"})
    item = ET.SubElement(strings, f"{{{MAIN_NS}}}si")
    ET.SubElement(item, f"{{{MAIN_NS}}}t").text = original
    parts["xl/sharedStrings.xml"] = ET.tostring(strings, encoding="utf-8", xml_declaration=True)
    content = ET.fromstring(parts["[Content_Types].xml"])
    ET.SubElement(content, f"{{{content_ns}}}Override", {
        "PartName": "/xl/sharedStrings.xml",
        "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
    })
    parts["[Content_Types].xml"] = ET.tostring(content, encoding="utf-8", xml_declaration=True)
    rels = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
    ET.SubElement(rels, f"{{{rel_ns}}}Relationship", {
        "Id": "rIdSharedSecurity", "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
        "Target": "sharedStrings.xml",
    })
    parts["xl/_rels/workbook.xml.rels"] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
    return rebuild(parts, infos)


def amount_scan(source: bytes):
    return DetectionCoordinator().scan(
        source, ["财务"], {SensitiveType.FINANCIAL_AMOUNT},
        [ColumnRule("财务", "B", SensitiveType.FINANCIAL_AMOUNT)],
        restrict_to_column_rules=True,
    )


class SecureAiExportTests(unittest.TestCase):
    def test_financial_amount_becomes_numeric_888_and_currency_format_stays(self) -> None:
        source = inject_formula_cache(financial_source())
        scan = amount_scan(source)
        self.assertEqual([(item.coordinate, item.original_kind) for item in scan.candidates], [("B2", "float")])
        result = ProcessingService().process(source=source, source_filename="finance.xlsx", scanned_sheets=["财务"], candidates=scan.candidates)
        workbook = load_workbook(BytesIO(result.output_bytes), data_only=False)
        try:
            cell = workbook["财务"]["B2"]
            self.assertEqual(cell.value, 888)
            self.assertIsInstance(cell.value, int)
            self.assertEqual(cell.number_format, '¥#,##0.00')
            self.assertTrue(cell.font.bold)
            self.assertEqual(workbook["财务"]["C2"].value, 42)
        finally:
            workbook.close()
        self.assertTrue(result.validation.passed)

    def test_percentage_date_boolean_blank_and_formula_are_not_amount_candidates(self) -> None:
        scan = amount_scan(financial_source())
        self.assertEqual([item.coordinate for item in scan.candidates], ["B2"])
        self.assertTrue(any("百分比" in warning for warning in scan.warnings))

    def test_formula_text_is_preserved_and_cached_real_total_is_removed(self) -> None:
        source = inject_formula_cache(financial_source())
        result = ProcessingService().process(source=source, source_filename="finance.xlsx", scanned_sheets=["财务"], candidates=amount_scan(source).candidates)
        workbook = load_workbook(BytesIO(result.output_bytes), data_only=False)
        try:
            self.assertEqual(workbook["财务"]["D2"].value, "=SUM(B2:B6)")
        finally:
            workbook.close()
        with ZipFile(BytesIO(result.output_bytes)) as archive:
            root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            formula = root.find(f".//{{{MAIN_NS}}}c[@r='D2']")
            self.assertIsNotNone(formula.find(f"{{{MAIN_NS}}}f"))
            self.assertIsNone(formula.find(f"{{{MAIN_NS}}}v"))
        self.assertGreaterEqual(result.formula_caches_cleared, 1)

    def test_text_amount_fragment_becomes_text_888(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "财务"
        sheet.append(["金额"])
        sheet.append(["合同金额为 123,456.78 元"])
        source = save(workbook)
        scan = DetectionCoordinator().scan(
            source, ["财务"], {SensitiveType.FINANCIAL_AMOUNT},
            [ColumnRule("财务", "A", SensitiveType.FINANCIAL_AMOUNT)],
            restrict_to_column_rules=True,
        )
        result = ProcessingService().process(source=source, source_filename="text-amount.xlsx", scanned_sheets=["财务"], candidates=scan.candidates)
        workbook = load_workbook(BytesIO(result.output_bytes))
        try:
            self.assertEqual(workbook["财务"]["A2"].value, "合同金额为 888 元")
        finally:
            workbook.close()

    def test_shared_string_original_is_removed(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "客户"
        sheet.append(["客户名称"])
        sheet.append(["北京甲公司"])
        source = to_shared_string(save(workbook), ["A2"])
        scan = DetectionCoordinator().scan(source, ["客户"], {SensitiveType.CUSTOMER}, [ColumnRule("客户", "A", SensitiveType.CUSTOMER)])
        result = ProcessingService().process(source=source, source_filename="shared.xlsx", scanned_sheets=["客户"], candidates=scan.candidates)
        with ZipFile(BytesIO(result.output_bytes)) as archive:
            self.assertNotIn("北京甲公司".encode(), archive.read("xl/sharedStrings.xml"))
        self.assertEqual(result.shared_strings_removed, 1)

    def test_shared_string_still_referenced_by_unselected_cell_blocks_export(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "客户"
        sheet.append(["客户名称"])
        sheet.append(["北京甲公司"])
        sheet.append(["北京甲公司"])
        source = to_shared_string(save(workbook), ["A2", "A3"])
        scan = DetectionCoordinator().scan(source, ["客户"], {SensitiveType.CUSTOMER}, [ColumnRule("客户", "A", SensitiveType.CUSTOMER)])
        scan.candidates[1].enabled = False
        with self.assertRaisesRegex(UnsafeWorkbookError, "A3"):
            ProcessingService().process(source=source, source_filename="shared.xlsx", scanned_sheets=["客户"], candidates=scan.candidates)

    def test_same_customer_across_sheets_uses_same_visible_marker(self) -> None:
        workbook = Workbook()
        first = workbook.active
        first.title = "A"
        first.append(["客户名称"]); first.append(["北京甲公司"])
        second = workbook.create_sheet("B")
        second.append(["客户名称"]); second.append(["北京甲公司"])
        source = save(workbook)
        rules = [ColumnRule("A", "A", SensitiveType.CUSTOMER), ColumnRule("B", "A", SensitiveType.CUSTOMER)]
        scan = DetectionCoordinator().scan(source, ["A", "B"], {SensitiveType.CUSTOMER}, rules)
        result = ProcessingService().process(source=source, source_filename="names.xlsx", scanned_sheets=["A", "B"], candidates=scan.candidates)
        workbook = load_workbook(BytesIO(result.output_bytes))
        try:
            self.assertEqual(workbook["A"]["A2"].value, "【已脱敏-客户-001】")
            self.assertEqual(workbook["A"]["A2"].value, workbook["B"]["A2"].value)
        finally:
            workbook.close()

    def test_hidden_sheet_not_scanned_blocks_export(self) -> None:
        workbook = Workbook()
        visible = workbook.active
        visible.title = "A"; visible.append(["客户名称"]); visible.append(["甲公司"])
        hidden = workbook.create_sheet("隐藏"); hidden.sheet_state = "hidden"; hidden.append(["备注"])
        source = save(workbook)
        scan = DetectionCoordinator().scan(source, ["A"], {SensitiveType.CUSTOMER}, [ColumnRule("A", "A", SensitiveType.CUSTOMER)])
        with self.assertRaisesRegex(UnsafeWorkbookError, "隐藏工作表"):
            ProcessingService().process(source=source, source_filename="hidden.xlsx", scanned_sheets=["A"], candidates=scan.candidates)

    def test_pivot_cache_blocks_export(self) -> None:
        self._assert_component_blocks("xl/pivotCache/pivotCacheDefinition1.xml")

    def test_chart_cache_blocks_export(self) -> None:
        self._assert_component_blocks("xl/charts/chart1.xml")

    def test_external_connection_blocks_export(self) -> None:
        self._assert_component_blocks("xl/externalLinks/externalLink1.xml")

    def test_unscanned_comment_blocks_export(self) -> None:
        workbook = Workbook(); sheet = workbook.active; sheet.title = "A"
        sheet.append(["客户名称"]); sheet.append(["甲公司"])
        sheet["B2"].comment = Comment("敏感备注", "tester")
        source = save(workbook)
        scan = DetectionCoordinator().scan(source, ["A"], {SensitiveType.CUSTOMER}, [ColumnRule("A", "A", SensitiveType.CUSTOMER)])
        with self.assertRaises(UnsafeWorkbookError):
            ProcessingService().process(source=source, source_filename="comment.xlsx", scanned_sheets=["A"], candidates=scan.candidates)

    def test_very_hidden_sheet_always_blocks_export(self) -> None:
        workbook = Workbook(); sheet = workbook.active; sheet.title = "A"
        sheet.append(["客户名称"]); sheet.append(["甲公司"])
        secret = workbook.create_sheet("秘密"); secret.sheet_state = "veryHidden"; secret.append(["备注"])
        source = save(workbook)
        scan = DetectionCoordinator().scan(source, ["A", "秘密"], {SensitiveType.CUSTOMER}, [ColumnRule("A", "A", SensitiveType.CUSTOMER)])
        with self.assertRaisesRegex(UnsafeWorkbookError, "veryHidden"):
            ProcessingService().process(source=source, source_filename="very-hidden.xlsx", scanned_sheets=["A", "秘密"], candidates=scan.candidates)

    def test_password_protected_workbook_blocks_export(self) -> None:
        workbook = Workbook(); sheet = workbook.active; sheet.title = "A"
        sheet.append(["客户名称"]); sheet.append(["甲公司"])
        workbook.security.lockStructure = True
        workbook.security.set_workbook_password("test")
        source = save(workbook)
        scan = DetectionCoordinator().scan(source, ["A"], {SensitiveType.CUSTOMER}, [ColumnRule("A", "A", SensitiveType.CUSTOMER)])
        with self.assertRaisesRegex(UnsafeWorkbookError, "保护"):
            ProcessingService().process(source=source, source_filename="protected.xlsx", scanned_sheets=["A"], candidates=scan.candidates)

    def _assert_component_blocks(self, name: str) -> None:
        workbook = Workbook(); sheet = workbook.active; sheet.title = "A"
        sheet.append(["客户名称"]); sheet.append(["甲公司"])
        source = add_part(save(workbook), name)
        scan = DetectionCoordinator().scan(source, ["A"], {SensitiveType.CUSTOMER}, [ColumnRule("A", "A", SensitiveType.CUSTOMER)])
        with self.assertRaises(UnsafeWorkbookError):
            ProcessingService().process(source=source, source_filename="risk.xlsx", scanned_sheets=["A"], candidates=scan.candidates)

    def test_source_sha_is_unchanged_and_output_reopens(self) -> None:
        source = financial_source(); digest = hashlib.sha256(source).hexdigest()
        result = ProcessingService().process(source=source, source_filename="finance.xlsx", scanned_sheets=["财务"], candidates=amount_scan(source).candidates)
        self.assertEqual(hashlib.sha256(source).hexdigest(), digest)
        workbook = load_workbook(BytesIO(result.output_bytes)); workbook.close()

    def test_report_contains_no_complete_sensitive_values(self) -> None:
        source = financial_source()
        result = ProcessingService().process(source=source, source_filename="finance.xlsx", scanned_sheets=["财务"], candidates=amount_scan(source).candidates)
        with ZipFile(BytesIO(result.report_bytes)) as archive:
            payload = b"".join(archive.read(name) for name in archive.namelist())
        self.assertNotIn(b"123456.78", payload)

    def test_validator_rejects_non_target_cell_change(self) -> None:
        source, operations, patch = self._simple_patch()
        workbook = load_workbook(BytesIO(patch.output_bytes)); workbook["A"]["B2"] = 999
        altered = save(workbook)
        validation = ExcelValidator().validate(source, altered, operations, set(ZipFile(BytesIO(source)).namelist()), patch.replacement_count)
        self.assertFalse(validation.passed)
        self.assertTrue(any(not check.passed and "非目标单元格" in check.name for check in validation.checks))

    def test_validator_rejects_non_allowlisted_ooxml_change(self) -> None:
        source, operations, patch = self._simple_patch()
        with ZipFile(BytesIO(patch.output_bytes)) as archive:
            parts = {item.filename: archive.read(item.filename) for item in archive.infolist()}
            infos = archive.infolist()
        parts["docProps/core.xml"] = parts["docProps/core.xml"].replace(b"</", b" \n</", 1)
        altered = rebuild(parts, infos)
        validation = ExcelValidator().validate(source, altered, operations, patch.modified_parts, patch.replacement_count)
        self.assertFalse(validation.passed)
        self.assertTrue(any(not check.passed and "OOXML" in check.name for check in validation.checks))

    def _simple_patch(self):
        workbook = Workbook(); sheet = workbook.active; sheet.title = "A"
        sheet.append(["客户名称", "数字"]); sheet.append(["甲公司", 5])
        source = save(workbook)
        scan = DetectionCoordinator().scan(source, ["A"], {SensitiveType.CUSTOMER}, [ColumnRule("A", "A", SensitiveType.CUSTOMER)])
        operations = ReplacementPlanner().build(scan.candidates)
        return source, operations, OoxmlPatcher().patch(source, operations)

    def test_numeric_and_text_888_hashes_are_distinct(self) -> None:
        self.assertNotEqual(value_hash(888), value_hash("888"))


if __name__ == "__main__":
    unittest.main()
